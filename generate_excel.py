import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('/home/user/C6/C6A_Dashboard/data/submissions.json','r') as f:
    data = json.load(f)

wb = Workbook()

# Styles
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="FF1E3A5F", end_color="FF1E3A5F", fill_type="solid")
green_fill = PatternFill(start_color="FFD1FAE5", end_color="FFD1FAE5", fill_type="solid")
yellow_fill = PatternFill(start_color="FFFEF3C7", end_color="FFFEF3C7", fill_type="solid")
red_fill = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
red_font = Font(color="991B1B", bold=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def style_header(ws, row, cols):
    for col in range(1, cols+1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border

# ========== Sheet 1: 总览 ==========
ws1 = wb.active
ws1.title = "总览"
cids = list(data['challenges'].keys())

# Title
ws1.merge_cells('A1:' + get_column_letter(len(cids)+2) + '1')
ws1['A1'] = "Elite 20 提交数据总览"
ws1['A1'].font = Font(bold=True, size=16)
ws1['A1'].alignment = center

# Header row
headers = ["学生"] + cids + ["完成率"]
for col, h in enumerate(headers, 1):
    ws1.cell(row=3, column=col, value=h)
style_header(ws1, 3, len(headers))

# Data rows
for row_idx, st in enumerate(data['students'], 4):
    ws1.cell(row=row_idx, column=1, value=st['name']).border = thin_border
    comp_count = 0
    for col_idx, cid in enumerate(cids, 2):
        sub = next((s for s in data['submissions'] if s['student']['name']==st['name'] and s['challenge']==cid), None)
        if sub:
            label = {"complete":"✅","partial":"⚠️","missing":"❌"}[sub['status']]
            fill = {"complete":green_fill,"partial":yellow_fill,"missing":red_fill}[sub['status']]
            if sub['status']=='complete': comp_count += 1
        else:
            label = "❌"; fill = red_fill
        cell = ws1.cell(row=row_idx, column=col_idx, value=label)
        cell.fill = fill; cell.alignment = center; cell.border = thin_border
    rate = round(comp_count/len(cids)*100)
    cell = ws1.cell(row=row_idx, column=len(cids)+2, value=f"{rate}%")
    cell.alignment = center; cell.border = thin_border
    cell.font = Font(bold=True, color="FF065F46" if rate>=70 else "FF92400E" if rate>=40 else "FF991B1B")

# Summary row
sum_row = len(data['students'])+4
ws1.cell(row=sum_row, column=1, value="挑战完成率").font = Font(bold=True)
for col_idx, cid in enumerate(cids, 2):
    subs = [s for s in data['submissions'] if s['challenge']==cid]
    comp = len([s for s in subs if s['status']=='complete'])
    rate = round(comp/len(subs)*100)
    cell = ws1.cell(row=sum_row, column=col_idx, value=f"{rate}%")
    cell.alignment = center; cell.font = Font(bold=True)

# Column widths
ws1.column_dimensions['A'].width = 12
for i in range(2, len(cids)+2):
    ws1.column_dimensions[get_column_letter(i)].width = 8
ws1.column_dimensions[get_column_letter(len(cids)+2)].width = 10

# ========== Sheet 2: 按学生 ==========
ws2 = wb.create_sheet("按学生")
headers2 = ["学生姓名","专业","挑战编号","挑战名称","提交时间","状态","完成度","缺失文件","达到等级","备注"]
for col, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=col, value=h)
style_header(ws2, 1, len(headers2))

row = 2
for st in data['students']:
    for sub in [s for s in data['submissions'] if s['student']['name']==st['name']]:
        status_label = {"complete":"完整","partial":"部分","missing":"缺失"}[sub['status']]
        vals = [st['name'], st['major'], sub['challenge'], sub['challenge_title'],
                sub['submitted_at'] or '', status_label, f"{sub['completeness']*100:.0f}%",
                ', '.join(sub['missing']), sub['level_achieved'], sub['notes']]
        for col, v in enumerate(vals, 1):
            cell = ws2.cell(row=row, column=col, value=v)
            cell.border = thin_border
            if sub['status']=='missing':
                cell.fill = red_fill
            elif sub['status']=='partial':
                cell.fill = yellow_fill
        row += 1

for col in range(1, len(headers2)+1):
    ws2.column_dimensions[get_column_letter(col)].width = 15

# ========== Sheet 3: 按挑战 ==========
ws3 = wb.create_sheet("按挑战")
headers3 = ["挑战编号","挑战名称","学生姓名","提交时间","状态","完成度","缺失文件","达到等级"]
for col, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=col, value=h)
style_header(ws3, 1, len(headers3))

row = 2
for cid, chal in data['challenges'].items():
    for sub in [s for s in data['submissions'] if s['challenge']==cid]:
        status_label = {"complete":"完整","partial":"部分","missing":"缺失"}[sub['status']]
        vals = [cid, chal['title'], sub['student']['name'], sub['submitted_at'] or '',
                status_label, f"{sub['completeness']*100:.0f}%", ', '.join(sub['missing']), sub['level_achieved']]
        for col, v in enumerate(vals, 1):
            cell = ws3.cell(row=row, column=col, value=v)
            cell.border = thin_border
            if sub['status']=='missing': cell.fill = red_fill
            elif sub['status']=='partial': cell.fill = yellow_fill
        row += 1

for col in range(1, len(headers3)+1):
    ws3.column_dimensions[get_column_letter(col)].width = 16

# ========== Sheet 4: 时间线 ==========
ws4 = wb.create_sheet("时间线")
headers4 = ["提交时间","学生姓名","挑战编号","挑战名称","状态","完成度"]
for col, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=col, value=h)
style_header(ws4, 1, len(headers4))

# Sort by submitted_at
timeline = [s for s in data['submissions'] if s['submitted_at']]
timeline.sort(key=lambda x: x['submitted_at'])

row = 2
for sub in timeline:
    status_label = {"complete":"完整","partial":"部分","missing":"缺失"}[sub['status']]
    vals = [sub['submitted_at'], sub['student']['name'], sub['challenge'], sub['challenge_title'],
            status_label, f"{sub['completeness']*100:.0f}%"]
    for col, v in enumerate(vals, 1):
        cell = ws4.cell(row=row, column=col, value=v)
        cell.border = thin_border
    row += 1

for col in range(1, len(headers4)+1):
    ws4.column_dimensions[get_column_letter(col)].width = 20

# ========== Sheet 5: 缺失项 ==========
ws5 = wb.create_sheet("缺失项")
headers5 = ["学生姓名","挑战编号","挑战名称","状态","缺失文件","提交时间"]
for col, h in enumerate(headers5, 1):
    ws5.cell(row=1, column=col, value=h)
style_header(ws5, 1, len(headers5))

row = 2
missing_items = [s for s in data['submissions'] if s['status'] != 'complete']
missing_items.sort(key=lambda x: (x['student']['name'], x['challenge']))
for sub in missing_items:
    status_label = {"partial":"⚠️ 部分提交","missing":"❌ 未提交"}[sub['status']]
    vals = [sub['student']['name'], sub['challenge'], sub['challenge_title'],
            status_label, ', '.join(sub['missing']), sub['submitted_at'] or '未提交']
    for col, v in enumerate(vals, 1):
        cell = ws5.cell(row=row, column=col, value=v)
        cell.border = thin_border
        cell.fill = red_fill if sub['status']=='missing' else yellow_fill
        if col == 1: cell.font = red_font
    row += 1

for col in range(1, len(headers5)+1):
    ws5.column_dimensions[get_column_letter(col)].width = 18

# Save
output_path = '/home/user/C6/C6A_Dashboard/report.xlsx'
wb.save(output_path)
print(f"Excel 报表已生成: {output_path}")
print(f"包含 {len(wb.sheetnames)} 个 Sheet: {', '.join(wb.sheetnames)}")
print(f"缺失项 Sheet 共 {len(missing_items)} 条记录")
